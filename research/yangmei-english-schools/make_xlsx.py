"""Build the Excel deliverable: 楊梅 20 km 名錄 + 全台英文補習班 + 統計."""
import json, os, statistics as st
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

SP = os.path.dirname(os.path.abspath(__file__))
rows = json.load(open(f"{SP}/english-schools-scale.json"))
rows.sort(key=lambda r: r["distance_km"])
tw = json.load(open(f"{SP}/tw-english-schools.json"))
OUT = f"{SP}/楊梅英語補習班名錄.xlsx"
VIEW = "https://afterschools.olc.tw/afterschools/view/"

F, INK = "Arial", "1F2A28"
HDR = PatternFill("solid", fgColor="0B5C50")
HDR2 = PatternFill("solid", fgColor="1F5F7A")
BAND_FILL = {"0–5 km": "D6E7E0", "5–10 km": "E4EFE9", "10–15 km": "EFF3EE", "15–20 km": "F7F8F6"}
TIER = {"大型": "8A4B12", "中型": "1F5F7A", "小型": "6B736E"}
LINK = Font(name=F, size=10, color="0B5C50", underline="single")
BORDER = Border(bottom=Side(style="thin", color="D7DCD4"))
BODY = Font(name=F, size=10, color=INK)

NUMFMT = {"distance_km": "0.00", "rating": "0.0", "reviews": "#,##0", "classrooms": "#,##0",
          "classroom_m2": "#,##0.0", "building_m2": "#,##0.0", "n_classes": "#,##0",
          "n_teachers": "#,##0", "approved_capacity": "#,##0",
          "approved_capacity_en": "#,##0", "lat": "0.000000", "lng": "0.000000"}

wb = Workbook()

# ══ 說明 ═══════════════════════════════════════════════════════════════
ws = wb.active
ws.title = "說明"
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 20
ws.column_dimensions["B"].width = 96


def line(r, k, v, size=10):
    ws.cell(r, 1, k).font = Font(name=F, size=size, bold=True, color=INK)
    c = ws.cell(r, 2, v)
    c.font = Font(name=F, size=size, color=INK)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    return r + 1


def sect(r, t, color="0B5C50"):
    ws.cell(r + 1, 1, t).font = Font(name=F, size=12, bold=True, color=color)
    return r + 2


n_mail = sum(1 for r in rows if r.get("email"))
tw_mail = sum(1 for r in tw if r["email"])
tw_gmail = sum(1 for r in tw if "gmail.com" in r["email"].lower())

ws["A1"] = "英語補習班名錄"
ws["A1"].font = Font(name=F, size=18, bold=True, color="0B5C50")
ws["A2"] = "楊梅永平路 20 公里圈 ＋ 全台立案英文補習班"
ws["A2"].font = Font(name=F, size=11, color="55625D")

r = 4
r = line(r, "擷取日期", "2026-09-09（Google 名單）／2026-09-10（立案資料與 email）")
r = sect(r, "分頁說明")
r = line(r, "楊梅20km名錄", f"{len(rows)} 家。以桃園市楊梅區永平路（24.920261, 121.178141）為圓心、"
                        f"半徑 20 公里的英文補習班。名單來自 Google Maps，含評分、官網、座標與距離；"
                        f"其中 {sum(1 for x in rows if x['reg_name'])} 家已對上立案資料。")
r = line(r, "全台英文補習班", f"{len(tw):,} 家。完全來自教育部立案資料，涵蓋全台 22 縣市。"
                        f"沒有 Google 評分、官網、座標與距離——那些需要 Places API，"
                        f"而免費配額每日僅 100 次，不可能涵蓋數千家。")
r = line(r, "規模統計", "整體、依距離級距、依規模分級、依行政區、全台依縣市。全部為公式，會自動重算。")
r = line(r, "未對上立案", f"{sum(1 for x in rows if not x['reg_name'])} 家 Google 找得到、"
                      "但對不上立案資料者。")

r = sect(r, "資料來源")
r = line(r, "名單／聯絡方式", "Google Maps Places API (New) places:searchNearby。六角網格 85 格覆蓋 "
                       "20 公里圓域，每格半徑 2,800 公尺；類型限 school 與 child_care_agency，"
                       "排除主類型為公立中小學與大學者。")
r = line(r, "立案資料／email／規模", "教育部「直轄市及各縣市短期補習班資訊管理系統」，經 "
                            "https://afterschools.olc.tw/ 取得。全台 17,836 家立案補習班中，"
                            f"外語類 4,819 家，篩出實際教英文且未廢止者 {len(tw):,} 家。")
r = line(r, "官方原始站", "https://bsb.kh.edu.tw/ 　立案資料權威版本，鏡站為二手；有出入以官方為準。")

r = sect(r, "email 欄位", "8A5A12")
r = line(r, "來源", "立案登記的聯絡信箱，非 Google 資料（Google Places 沒有 email 欄位）。")
r = line(r, "覆蓋率", f"楊梅名錄 {n_mail}/{len(rows)} 家；全台 {tw_mail:,}/{len(tw):,} 家（{tw_mail*100//len(tw)}%）。"
                    f"全台其中 {tw_gmail:,} 家是 Gmail。空白代表立案時未填或未更新。")
r = line(r, "注意", "這是負責人或班務的登記信箱，未必是招生窗口，也可能已停用。"
                 "群發前請先小量測試退信率，並遵守個資法與商業電子郵件相關規範。")

r = sect(r, "重要限制", "8A5A12")
for k, v in [
    ("沒有學生人數", "實際在學人數不屬於公開資訊，兩個來源都沒有這個欄位。"),
    ("核定容量的算法", "各核准科目「班級數 × 每班核准人數」加總，同一批學生修多科會重複計算，"
                 "會出現「4 間教室卻核定 1,773 人」。判斷規模請看教室數與班舍總面積。"),
    ("楊梅名單不完整", "Places API 單次上限 20 筆，85 格中有 36 格觸頂，中壢、桃園、八德等密集區有遺漏。"
                  "同範圍立案外語類實有 1,064 家，遠多於 Google 能看到的 159 家。"),
    ("全台名單較完整但無距離", "立案資料是普查等級，但沒有經緯度。要算距離需 Geocoding API。"),
    ("英文判定方式", "全台分頁以「核准科目含英文或美語」為準，其次看名稱；"
                "純日語、韓語補習班已剔除，已廢止／註銷者亦剔除。"),
]:
    r = line(r, k, v)

r = sect(r, "規模分級標準")
for k, v in [("大型", "班舍總面積 ≥ 400 ㎡ 或教室 ≥ 8 間"), ("中型", "≥ 200 ㎡ 或 ≥ 5 間"),
             ("小型", "其餘"), ("空白", "未對上立案資料，無規模數據")]:
    r = line(r, k, v)


# ══ 共用寫表函式 ════════════════════════════════════════════════════════
def write_sheet(ws, cols, data, fill=HDR, band_col=None, freeze="A2"):
    for i, (title, _, w) in enumerate(cols, 1):
        c = ws.cell(1, i, title)
        c.font = Font(name=F, size=10, bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 30
    for ri, rec in enumerate(data, 2):
        for ci, (title, key, _) in enumerate(cols, 1):
            if key == "_reg_url":
                ids = (rec.get("reg_id") or "").split()
                c = ws.cell(ri, ci, "立案頁" if ids else None)
                if ids:
                    c.hyperlink, c.font = VIEW + ids[0], LINK
            else:
                v = rec.get(key, "")
                c = ws.cell(ri, ci, v if v != "" else None)
                if key in ("website", "maps") and v:
                    c.hyperlink, c.value = v, {"website": "官網", "maps": "地圖"}[key]
                    c.font = LINK
                elif key == "size_tier" and v:
                    c.font = Font(name=F, size=10, bold=True, color=TIER[v])
                    c.alignment = Alignment(horizontal="center")
                else:
                    c.font = BODY
            if key in NUMFMT:
                c.number_format = NUMFMT[key]
            if key == band_col:
                c.fill = PatternFill("solid", fgColor=BAND_FILL[rec[key]])
                c.alignment = Alignment(horizontal="center")
            if key in ("name", "reg_name", "address", "category", "en_subjects"):
                c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = BORDER
    ws.freeze_panes = freeze
    ws.sheet_view.showGridLines = False


# ══ 楊梅20km名錄 ═══════════════════════════════════════════════════════
COLS = [
    ("距離級距", "band", 11), ("距離(km)", "distance_km", 10),
    ("名稱（Google 招牌）", "name", 40), ("立案名稱", "reg_name", 38),
    ("縣市", "city", 9), ("行政區", "district", 9), ("地址", "address", 32),
    ("電話", "phone", 15), ("Email", "email", 28), ("官網", "website", 10),
    ("評分", "rating", 7), ("評論數", "reviews", 8),
    ("規模", "size_tier", 7), ("教室數", "classrooms", 8),
    ("教室面積(㎡)", "classroom_m2", 12), ("班舍總面積(㎡)", "building_m2", 14),
    ("核准班級數", "n_classes", 11), ("登記師資(人)", "n_teachers", 12),
    ("核定容量(人)", "approved_capacity", 12), ("英文科核定(人)", "approved_capacity_en", 14),
    ("立案類別", "category", 20), ("立案日期", "licensed_on", 11),
    ("負責人", "principal", 11), ("班主任", "director", 11),
    ("比對方式", "match", 9), ("立案頁", "_reg_url", 10), ("Google Maps", "maps", 12),
    ("緯度", "lat", 11), ("經度", "lng", 11),
]
ym = wb.create_sheet("楊梅20km名錄")
write_sheet(ym, COLS, rows, band_col="band", freeze="D2")
t = Table(displayName="楊梅名錄", ref=f"A1:{get_column_letter(len(COLS))}{len(rows)+1}")
t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
ym.add_table(t)

# ══ 全台英文補習班 ══════════════════════════════════════════════════════
TWC = [
    ("縣市", "city", 10), ("行政區", "district", 10), ("立案名稱", "reg_name", 42),
    ("地址", "address", 36), ("電話", "phone", 15), ("Email", "email", 30),
    ("規模", "size_tier", 7), ("教室數", "classrooms", 8),
    ("教室面積(㎡)", "classroom_m2", 12), ("班舍總面積(㎡)", "building_m2", 14),
    ("核准班級數", "n_classes", 11), ("登記師資(人)", "n_teachers", 12),
    ("核定容量(人)", "approved_capacity", 12), ("英文科核定(人)", "approved_capacity_en", 14),
    ("英文相關核准科目", "en_subjects", 22), ("立案類別", "category", 22),
    ("立案情形", "reg_status", 10), ("立案日期", "licensed_on", 11),
    ("負責人", "principal", 11), ("班主任", "director", 11), ("立案頁", "_reg_url", 10),
]
tws = wb.create_sheet("全台英文補習班")
write_sheet(tws, TWC, tw, fill=HDR2, freeze="D2")
t2 = Table(displayName="全台名錄", ref=f"A1:{get_column_letter(len(TWC))}{len(tw)+1}")
t2.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
tws.add_table(t2)

# ══ 規模統計 ═══════════════════════════════════════════════════════════
sm = wb.create_sheet("規模統計")
sm.sheet_view.showGridLines = False
for col, w in zip("ABCDEFGH", (16, 11, 12, 13, 15, 13, 13, 13)):
    sm.column_dimensions[col].width = w
N, M = len(rows) + 1, len(tw) + 1
Y = "楊梅20km名錄!"
BAND, TIER_C = f"{Y}$A$2:$A${N}", f"{Y}$M$2:$M${N}"
ROOM, AREA = f"{Y}$N$2:$N${N}", f"{Y}$P$2:$P${N}"
TCH, REGN = f"{Y}$R$2:$R${N}", f"{Y}$D$2:$D${N}"
DIST, RATE, MAIL = f"{Y}$F$2:$F${N}", f"{Y}$K$2:$K${N}", f"{Y}$I$2:$I${N}"
T = "全台英文補習班!"
TCITY, TTIER = f"{T}$A$2:$A${M}", f"{T}$G$2:$G${M}"
TROOM, TAREA = f"{T}$H$2:$H${M}", f"{T}$J$2:$J${M}"
TTCH, TMAIL = f"{T}$L$2:$L${M}", f"{T}$F$2:$F${M}"


def head(r, t):
    sm.cell(r, 1, t).font = Font(name=F, size=12, bold=True, color="0B5C50")
    return r + 1


def hrow(r, labels, fill=HDR):
    for i, t in enumerate(labels, 1):
        c = sm.cell(r, i, t)
        c.font = Font(name=F, size=10, bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    return r + 1


def put(r, col, formula, fmt="#,##0"):
    c = sm[f"{col}{r}"]
    c.value, c.font, c.number_format = formula, BODY, fmt
    return c


sm["A1"] = "規模統計"
sm["A1"].font = Font(name=F, size=16, bold=True, color="0B5C50")
sm["A2"] = "所有數字皆為公式，改動名錄分頁會自動重算。"
sm["A2"].font = Font(name=F, size=10, color="55625D")

r = head(4, "楊梅 20 公里圈 — 整體")
r = hrow(r, ["", "家數", "有立案資料", "有 Email", "教室數中位數", "班舍面積中位數(㎡)",
             "師資中位數(人)", "平均班舍面積(㎡)"])
sm.cell(r, 1, "全部").font = Font(name=F, size=10, bold=True, color=INK)
for col, f, fmt in [("B", f"=COUNTA({BAND})", "#,##0"), ("C", f"=COUNTA({REGN})", "#,##0"),
                    ("D", f'=COUNTIF({MAIL},"?*")', "#,##0"),
                    ("E", f"=MEDIAN({ROOM})", "#,##0.#"), ("F", f"=MEDIAN({AREA})", "#,##0.0"),
                    ("G", f"=MEDIAN({TCH})", "#,##0.#"), ("H", f"=AVERAGE({AREA})", "#,##0.0")]:
    put(r, col, f, fmt)
r += 2

r = head(r, "楊梅 20 公里圈 — 依距離級距")
r = hrow(r, ["距離級距", "家數", "有立案資料", "有 Email", "平均教室數",
             "平均班舍面積(㎡)", "平均師資(人)", "大型"])
for band in ["0–5 km", "5–10 km", "10–15 km", "15–20 km"]:
    c0 = sm.cell(r, 1, band)
    c0.font, c0.fill = BODY, PatternFill("solid", fgColor=BAND_FILL[band])
    put(r, "B", f"=COUNTIF({BAND},$A{r})")
    put(r, "C", f'=COUNTIFS({BAND},$A{r},{REGN},"<>")')
    put(r, "D", f'=COUNTIFS({BAND},$A{r},{MAIL},"?*")')
    put(r, "E", f'=IFERROR(AVERAGEIFS({ROOM},{BAND},$A{r},{ROOM},">0"),"")', "0.0")
    put(r, "F", f'=IFERROR(AVERAGEIFS({AREA},{BAND},$A{r},{AREA},">0"),"")', "0.0")
    put(r, "G", f'=IFERROR(AVERAGEIFS({TCH},{BAND},$A{r},{TCH},">0"),"")', "0.0")
    put(r, "H", f'=COUNTIFS({BAND},$A{r},{TIER_C},"大型")')
    r += 1
r += 1

r = head(r, "楊梅 20 公里圈 — 依規模分級")
r = hrow(r, ["規模", "家數", "平均教室數", "平均班舍面積(㎡)", "平均師資(人)",
             "平均核准班級數", "平均評分", "占比"])
first = r
for tier in ["大型", "中型", "小型"]:
    sm.cell(r, 1, tier).font = Font(name=F, size=10, bold=True, color=TIER[tier])
    put(r, "B", f"=COUNTIF({TIER_C},$A{r})")
    put(r, "C", f'=IFERROR(AVERAGEIFS({ROOM},{TIER_C},$A{r}),"")', "0.0")
    put(r, "D", f'=IFERROR(AVERAGEIFS({AREA},{TIER_C},$A{r}),"")', "0.0")
    put(r, "E", f'=IFERROR(AVERAGEIFS({TCH},{TIER_C},$A{r}),"")', "0.0")
    put(r, "F", f'=IFERROR(AVERAGEIFS({Y}$Q$2:$Q${N},{TIER_C},$A{r}),"")', "0.0")
    put(r, "G", f'=IFERROR(AVERAGEIFS({RATE},{TIER_C},$A{r},{RATE},">0"),"")', "0.0")
    put(r, "H", f'=IFERROR(B{r}/SUM($B${first}:$B${first+2}),"")', "0.0%")
    r += 1
r += 1

r = head(r, "楊梅 20 公里圈 — 依行政區")
r = hrow(r, ["行政區", "家數", "有立案資料", "有 Email", "平均教室數",
             "平均班舍面積(㎡)", "平均師資(人)", "大型"])
for d, _ in Counter(x["district"] for x in rows).most_common():
    sm.cell(r, 1, d).font = BODY
    put(r, "B", f"=COUNTIF({DIST},$A{r})")
    put(r, "C", f'=COUNTIFS({DIST},$A{r},{REGN},"<>")')
    put(r, "D", f'=COUNTIFS({DIST},$A{r},{MAIL},"?*")')
    put(r, "E", f'=IFERROR(AVERAGEIFS({ROOM},{DIST},$A{r},{ROOM},">0"),"")', "0.0")
    put(r, "F", f'=IFERROR(AVERAGEIFS({AREA},{DIST},$A{r},{AREA},">0"),"")', "0.0")
    put(r, "G", f'=IFERROR(AVERAGEIFS({TCH},{DIST},$A{r},{TCH},">0"),"")', "0.0")
    put(r, "H", f'=COUNTIFS({DIST},$A{r},{TIER_C},"大型")')
    r += 1
r += 1

r = head(r, "全台英文補習班 — 依縣市")
r = hrow(r, ["縣市", "家數", "有 Email", "平均教室數", "平均班舍面積(㎡)",
             "平均師資(人)", "大型", "小型"], fill=HDR2)
tw_first = r
for c, _ in Counter(x["city"] or "（不明）" for x in tw).most_common():
    sm.cell(r, 1, c).font = BODY
    put(r, "B", f"=COUNTIF({TCITY},$A{r})")
    put(r, "C", f'=COUNTIFS({TCITY},$A{r},{TMAIL},"?*")')
    put(r, "D", f'=IFERROR(AVERAGEIFS({TROOM},{TCITY},$A{r},{TROOM},">0"),"")', "0.0")
    put(r, "E", f'=IFERROR(AVERAGEIFS({TAREA},{TCITY},$A{r},{TAREA},">0"),"")', "0.0")
    put(r, "F", f'=IFERROR(AVERAGEIFS({TTCH},{TCITY},$A{r},{TTCH},">0"),"")', "0.0")
    put(r, "G", f'=COUNTIFS({TCITY},$A{r},{TTIER},"大型")')
    put(r, "H", f'=COUNTIFS({TCITY},$A{r},{TTIER},"小型")')
    r += 1
sm.cell(r, 1, "合計").font = Font(name=F, size=10, bold=True, color=INK)
put(r, "B", f"=SUM(B{tw_first}:B{r-1})")
put(r, "C", f"=SUM(C{tw_first}:C{r-1})")
put(r, "G", f"=SUM(G{tw_first}:G{r-1})")
put(r, "H", f"=SUM(H{tw_first}:H{r-1})")

sm.cell(r + 2, 1, "註：平均值只計有數據者；未對上立案資料的補習班不列入教室／面積／師資的平均。") \
  .font = Font(name=F, size=9, italic=True, color="7D8A84")

# ══ 未對上立案 ═════════════════════════════════════════════════════════
un = wb.create_sheet("未對上立案")
UC = [("距離(km)", "distance_km", 10), ("名稱", "name", 46), ("行政區", "district", 10),
      ("地址", "address", 36), ("電話", "phone", 15), ("評分", "rating", 7),
      ("評論數", "reviews", 8), ("Google Maps", "maps", 12)]
miss = [x for x in rows if not x["reg_name"]]
write_sheet(un, UC, miss, fill=PatternFill("solid", fgColor="8A5A12"))
un.cell(len(miss) + 3, 1,
        "Google 找得到、但對不上教育部立案資料者。多為登記在「兒童課後照顧服務中心」名下"
        "（另一套系統），或 Google 地址與立案地址不一致。"
        "要確認立案狀態請到 https://bsb.kh.edu.tw/ 以名稱或地址查詢。") \
  .font = Font(name=F, size=9, italic=True, color="7D8A84")

wb.save(OUT)
print(f"saved {OUT}")
print(f"楊梅20km名錄 {len(rows)} 列（email {n_mail}）｜全台英文補習班 {len(tw):,} 列"
      f"（email {tw_mail:,}，Gmail {tw_gmail:,}）｜未對上立案 {len(miss)} 列")
